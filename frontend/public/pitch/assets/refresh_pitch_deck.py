"""
refresh_pitch_deck.py — inject the latest UI screenshots into the deck.
Slide 4: pyramid graph hero. Slide 9: graph / chat / impact thumbs.
Then export pitch-deck.pdf via PowerPoint COM and rebuild thumbnail.png.
Run from anywhere: paths are anchored to this file's directory.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFont
import os

HERE   = os.path.dirname(os.path.abspath(__file__))
SHOTS  = os.path.normpath(os.path.join(HERE, "..", "..", "..", "..", "docs", "screenshots"))
DECK   = os.path.normpath(os.path.join(HERE, "..", "pitch-deck.pptx"))
PDF    = os.path.normpath(os.path.join(HERE, "..", "pitch-deck.pdf"))
THUMB  = os.path.normpath(os.path.join(HERE, "..", "thumbnail.png"))

def shot(name): return os.path.join(SHOTS, name)
def asset(name): return os.path.join(HERE, name)

# ── 1. Crop fresh screenshots to frame ratios ────────────────────────────────
def crop_to(src, box, out, target_ratio):
    im = Image.open(src).crop(box)
    w, h = im.size
    r = w / h
    if r > target_ratio:                     # too wide → trim sides (centered)
        nw = int(h * target_ratio)
        x0 = (w - nw) // 2
        im = im.crop((x0, 0, x0 + nw, h))
    elif r < target_ratio:                   # too tall → trim bottom-biased
        nh = int(w / target_ratio)
        im = im.crop((0, 0, w, nh))
    im.save(out)
    return im.size

# Slide 4 hero — frame ratio 6.83 / 3.00 ≈ 2.277 (full pyramid band)
crop_to(shot("graph.png"), (300, 190, 1440, 690), asset("pitch-assets-hero.png"), 6.83 / 3.00)
# Slide 9 thumbs — frame ratio 2.40 / 0.89 ≈ 2.697
crop_to(shot("graph.png"),  (330, 170, 1410, 570), asset("pitch-assets-thumb-graph.png"),  2.697)
crop_to(shot("chat.png"),   (300, 380, 1440, 803), asset("pitch-assets-thumb-chat.png"),   2.697)
crop_to(shot("impact.png"), (300,  56, 1400, 464), asset("pitch-assets-thumb-impact.png"), 2.697)
print("crops OK")

# ── 2. Inject into the deck ──────────────────────────────────────────────────
prs = Presentation(DECK)

def delete_shape(sh):
    sh._element.getparent().remove(sh._element)

def add_framed_pic(slide, img, l, t, w, h):
    pic = slide.shapes.add_picture(img, Inches(l), Inches(t), Inches(w), Inches(h))
    pic.line.color.rgb = RGBColor(0x14, 0x14, 0x14)
    pic.line.width = Pt(1.5)
    return pic

# Slide 4 — replace the single screenshot picture in place
s4 = prs.slides[3]
replaced = 0
for sh in list(s4.shapes):
    if sh.shape_type == 13:  # PICTURE
        l, t, w, h = (sh.left.inches, sh.top.inches, sh.width.inches, sh.height.inches)
        delete_shape(sh)
        add_framed_pic(s4, asset("pitch-assets-hero.png"), l, t, w, h)
        replaced += 1
assert replaced == 1, f"slide 4: expected 1 picture, replaced {replaced}"

# Slide 9 — replace the three 2.40×0.89 row thumbs (top→bottom = graph/chat/impact)
s9 = prs.slides[8]
rows = []
for sh in list(s9.shapes):
    if sh.shape_type == 13 and abs(sh.width.inches - 2.40) < 0.05 and abs(sh.height.inches - 0.89) < 0.05:
        rows.append(sh)
rows.sort(key=lambda s: s.top.inches)
assert len(rows) == 3, f"slide 9: expected 3 thumbs, found {len(rows)}"
thumbs = ["pitch-assets-thumb-graph.png", "pitch-assets-thumb-chat.png", "pitch-assets-thumb-impact.png"]
for sh, img in zip(rows, thumbs):
    l, t, w, h = (sh.left.inches, sh.top.inches, sh.width.inches, sh.height.inches)
    delete_shape(sh)
    add_framed_pic(s9, asset(img), l, t, w, h)

prs.save(DECK)
print("deck updated →", DECK)

# ── 3. PPTX → PDF via PowerPoint COM ─────────────────────────────────────────
import win32com.client
app = win32com.client.Dispatch("PowerPoint.Application")
pres = app.Presentations.Open(DECK, ReadOnly=True, WithWindow=False)
pres.SaveAs(PDF, 32)  # 32 = ppSaveAsPDF
pres.Close()
app.Quit()
print("pdf exported →", PDF)

# ── 4. Rebuild thumbnail.png (1280×720, live-demo banner) ────────────────────
base = Image.open(shot("hero.png")).convert("RGB")
# center-crop to 16:9
w, h = base.size
nh = int(w * 9 / 16)
top = max(0, (h - nh) // 3)          # slight top bias — keeps the pyramid
base = base.crop((0, top, w, top + nh)).resize((1280, 720), Image.LANCZOS)

draw = ImageDraw.Draw(base)
try:
    font = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 34)
except OSError:
    font = ImageFont.load_default()
text = "RepoScope — live demo"
bbox = draw.textbbox((0, 0), text, font=font)
tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
# size the banner from the text so it never overflows
pad_x, pad_y = 28, 20
banner_w, banner_h = tw + pad_x * 2, th + pad_y * 2
margin = 28
x0, y0 = margin, 720 - banner_h - margin
draw.rounded_rectangle((x0, y0, x0 + banner_w, y0 + banner_h), radius=14, fill=(26, 26, 26))
draw.text((x0 + pad_x - bbox[0], y0 + pad_y - bbox[1]), text, font=font, fill=(255, 255, 255))
base.save(THUMB)
print("thumbnail →", THUMB)
