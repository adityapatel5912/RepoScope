"""
edit_pitch_deck.py — one-shot deck asset injection.
Slides: 1 (logo icon + footer), 4 (hero screenshot), 8 (session graphic +
attribution), 9 (three real feature screenshots). Idempotent-ish: run once.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from PIL import Image
import os

ROOT = os.path.dirname(os.path.abspath(__file__))

# ── 1. Prepare cropped images ────────────────────────────────────────────────
def crop_to(src, box, out, target_ratio=None):
    im = Image.open(os.path.join(ROOT, src))
    im = im.crop(box)
    if target_ratio:
        w, h = im.size
        r = w / h
        if r > target_ratio:            # too wide → trim sides
            nw = int(h * target_ratio)
            x0 = (w - nw) // 2
            im = im.crop((x0, 0, x0 + nw, h))
        elif r < target_ratio:          # too tall → trim bottom-biased
            nh = int(w / target_ratio)
            im = im.crop((0, 0, w, nh))
    im.save(os.path.join(ROOT, out))
    return im.size

# Slide 4 hero — frame inner ratio 6.83 / 3.00
crop_to("fixed-graph-v2.png", (0, 56, 1440, 688), "pitch-assets-hero.png", 6.83 / 3.00)
# Slide 9 thumbs — ratio 1.42 / 0.89
crop_to("fixed-graph-v2.png",    (320, 56, 1120, 560),  "pitch-assets-thumb-graph.png", 1.6)
crop_to("regress-graph-chat.png",(660, 420, 1428, 900), "pitch-assets-thumb-chat.png",  1.6)
crop_to("regress-impact.png",    (0,   56,  936, 640),  "pitch-assets-thumb-impact.png", 1.6)

# ── 2. Edit the deck ─────────────────────────────────────────────────────────
DECK = os.path.join(ROOT, "RepoScope — Pitch Deck.pptx")
prs = Presentation(DECK)

def delete_shape(sh):
    sh._element.getparent().remove(sh._element)

def add_framed_pic(slide, img, l, t, w, h, border=True):
    pic = slide.shapes.add_picture(os.path.join(ROOT, img), Inches(l), Inches(t), Inches(w), Inches(h))
    if border:
        pic.line.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0x14, 0x14, 0x14)
        pic.line.width = Pt(1.5)
    return pic

# ── Slide 1: logo icon + footer attribution ──
s1 = prs.slides[0]
for sh in list(s1.shapes):
    if sh.shape_type == 13 and abs(sh.width.inches - 0.62) < 0.05:
        l, t = sh.left, sh.top
        delete_shape(sh)
        s1.shapes.add_picture(os.path.join(ROOT, "pitch-assets-logo-icon.png"),
                              l - Inches(0.012), t, height=Inches(0.62))
for sh in s1.shapes:
    if sh.has_text_frame and "Built with IBM Bob" in sh.text_frame.text:
        sh.text_frame.text = "Built with IBM Bob + Z Code · IBM Bob 2.0 Hackathon · lablab.ai · 2026"

# ── Slide 4: hero screenshot into the placeholder frame ──
s4 = prs.slides[3]
kill, frame = [], None
for sh in s4.shapes:
    txt = sh.text_frame.text if sh.has_text_frame else ""
    if "SCREENSHOT PLACEHOLDER" in txt or "replace with a real" in txt or "RepoScope app — architecture graph" in txt:
        kill.append(sh)
    if sh.shape_type == 1 and abs(sh.width.inches - 6.87) < 0.1 and abs(sh.top.inches - 1.11) < 0.1:
        frame = sh
for sh in kill:
    delete_shape(sh)
add_framed_pic(s4, "pitch-assets-hero.png", 1.59, 1.13, 6.83, 3.00)

# ── Slide 8: session graphic + honest attribution ──
s8 = prs.slides[7]
kill = []
proof_frame = None
for sh in s8.shapes:
    txt = sh.text_frame.text if sh.has_text_frame else ""
    if "SCREENSHOT PLACEHOLDER" in txt or "session export" in txt or "IBM Bob session summary" in txt:
        kill.append(sh)
    if sh.shape_type == 1 and abs(sh.width.inches - 3.66) < 0.1 and abs(sh.top.inches - 1.63) < 0.1:
        proof_frame = sh
    if "Every file in RepoScope was written" in txt:
        sh.text_frame.text = ("Most tasks ran through IBM Bob; Z Code (GLM) paired on the "
                              "complex ones — graph layout, hardening, and this deck.")
    if txt.startswith("Built entirely with IBM Bob"):
        sh.text_frame.text = "Built with IBM Bob + Z Code."
for sh in kill:
    delete_shape(sh)
add_framed_pic(s8, "pitch-assets-session.png", 5.63, 1.65, 3.62, 2.30)

# ── Slide 9: three real screenshots in the demo rows ──
s9 = prs.slides[8]
kills, rows = [], []
for sh in s9.shapes:
    txt = sh.text_frame.text if sh.has_text_frame else ""
    if "SCREENSHOT PLACEHOLDER" in txt or "drop GIF or still" in txt:
        kills.append(sh)
    if sh.shape_type == 1 and abs(sh.width.inches - 9.16) < 0.1 and abs(sh.height.inches - 1.05) < 0.05:
        rows.append(sh)
for sh in kills:
    delete_shape(sh)
rows.sort(key=lambda s: s.top.inches)
thumbs = ["pitch-assets-thumb-graph.png", "pitch-assets-thumb-chat.png", "pitch-assets-thumb-impact.png"]
for row, img in zip(rows, thumbs):
    top = row.top.inches + (1.05 - 0.89) / 2
    add_framed_pic(s9, img, 1.14, top, 1.42, 0.89)

prs.save(DECK)
print("deck edited OK →", DECK)
